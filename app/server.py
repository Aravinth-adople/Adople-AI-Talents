import os
import json
import shutil
import csv
import io
import mimetypes
from pathlib import Path
from typing import List, Optional, Dict, Any
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Query, BackgroundTasks, Request
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from app.default_jds import DEFAULT_JDS, get_jd_by_id
from app.database import (
    init_db,
    get_candidates,
    get_candidate_by_id,
    update_candidate_status,
    update_candidate_ai_evaluation,
    get_statistics,
    upsert_candidate,
    delete_candidate,
    get_connection
)
from app.parser import parse_resume, extract_text_from_file
from app.classifier import classify_resume, ALL_CATEGORIES
from app.scorer import score_resume_against_jd, score_all_default_jds
from app.azure_ai import analyze_resume_with_ai
from app.batch_evaluator import BatchEvaluator

app = FastAPI(title="Executive Resume Classification & JD Match System")

# Enable CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Directories
BASE_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = BASE_DIR.parent
STATIC_DIR = BASE_DIR / "static"
TEMPLATES_DIR = BASE_DIR / "templates"
FRONTEND_DIST = PROJECT_ROOT / "frontend" / "dist"
UPLOADS_DIR = BASE_DIR / "uploads"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

# Mount static files
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")
if (FRONTEND_DIST / "assets").exists():
    app.mount("/assets", StaticFiles(directory=str(FRONTEND_DIST / "assets")), name="assets")

@app.get("/adople-ai-logo.png")
async def get_adople_logo():
    logo_path = STATIC_DIR / "adople-ai-logo.png"
    if not logo_path.exists():
        logo_path = FRONTEND_DIST / "adople-ai-logo.png"
    return FileResponse(str(logo_path), media_type="image/png")

@app.on_event("startup")
def on_startup():
    init_db()

# Models
class StatusUpdateRequest(BaseModel):
    status: str

class CustomJDRequest(BaseModel):
    id: str
    title: str
    department: Optional[str] = "Engineering"
    preferred_categories: List[str] = []
    key_skills: Dict[str, float] = {}
    project_keywords: List[str] = []
    raw_text: str

# ----------------- Frontend Page -----------------
@app.get("/", response_class=HTMLResponse)
async def serve_index():
    react_index = FRONTEND_DIST / "index.html"
    if react_index.exists():
        with open(react_index, "r", encoding="utf-8") as f:
            return HTMLResponse(f.read())
    return HTMLResponse("<h1>Loading React Multi-Page Platform...</h1>", status_code=200)


# ----------------- API Endpoints -----------------

@app.get("/api/jds")
async def list_jds():
    """List all available JDs (Default and Custom)."""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM custom_jds ORDER BY is_default DESC, title ASC")
    rows = cursor.fetchall()
    conn.close()
    
    jds = []
    for r in rows:
        jds.append({
            "id": r["id"],
            "title": r["title"],
            "department": r["department"],
            "badge_color": r["badge_color"],
            "preferred_categories": json.loads(r["preferred_categories"] or "[]"),
            "key_skills": json.loads(r["key_skills"] or "{}"),
            "project_keywords": json.loads(r["project_keywords"] or "[]"),
            "raw_text": r["raw_text"],
            "is_default": bool(r["is_default"])
        })
    return {"jds": jds}

@app.post("/api/jds")
async def create_custom_jd(jd: CustomJDRequest):
    """Add or update a custom Job Description."""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
    INSERT OR REPLACE INTO custom_jds (
        id, title, department, badge_color, preferred_categories,
        key_skills, project_keywords, raw_text, is_default
    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0)
    """, (
        jd.id,
        jd.title,
        jd.department,
        "#ec4899", # Pink badge for custom
        json.dumps(jd.preferred_categories),
        json.dumps(jd.key_skills),
        json.dumps(jd.project_keywords),
        jd.raw_text
    ))
    conn.commit()
    conn.close()
    return {"status": "success", "message": f"JD '{jd.title}' created."}

@app.get("/api/candidates")
async def fetch_candidates(
    search: str = Query("", description="Search by name, email, skills, projects, filename"),
    category: str = Query("", description="Job Category"),
    passout_year: Optional[str] = Query(None, description="Passout Year, including 2027"),
    min_cgpa: Optional[str] = Query(None, description="Minimum CGPA filter"),
    min_score: Optional[str] = Query(None, description="Minimum JD Match Score"),
    status: str = Query("", description="Candidate Status"),
    active_jd_id: str = Query("ai-engineer", description="Active JD for scoring"),
    sort_by: str = Query("score", description="Sort by score, cgpa, passout_year, name"),
    sort_order: str = Query("desc", description="Sort order: asc, desc"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=100)
):
    """Fetch paginated, filtered, sorted candidates."""
    py_val = None
    if passout_year and str(passout_year).strip():
        try:
            py_val = int(passout_year)
        except ValueError:
            pass

    cgpa_val = None
    if min_cgpa and str(min_cgpa).strip():
        try:
            cgpa_val = float(min_cgpa)
        except ValueError:
            pass

    score_val = None
    if min_score and str(min_score).strip():
        try:
            score_val = float(min_score)
        except ValueError:
            pass

    data = get_candidates(
        search=search,
        category=category,
        passout_year=py_val,
        min_cgpa=cgpa_val,
        min_score=score_val,
        status=status,
        active_jd_id=active_jd_id,
        sort_by=sort_by,
        sort_order=sort_order,
        page=page,
        page_size=page_size
    )
    return data

@app.get("/api/candidates/{candidate_id}")
async def get_candidate_details(candidate_id: int, jd_id: str = "ai-engineer"):
    """Get full details for a candidate with active JD breakdown."""
    c = get_candidate_by_id(candidate_id, active_jd_id=jd_id)
    if not c:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    jd = get_jd_by_id(jd_id)
    jd_analysis = score_resume_against_jd(c, jd)
    
    # Synchronize score so candidate.active_score and jd_analysis.score are always aligned
    score_val = c.get("active_score")
    if score_val is not None and score_val > 0:
        jd_analysis["score"] = score_val
    else:
        c["active_score"] = jd_analysis.get("score", 0.0)
    
    return {
        "candidate": c,
        "jd_analysis": jd_analysis
    }

@app.post("/api/candidates/{candidate_id}/status")
async def set_candidate_status(candidate_id: int, req: StatusUpdateRequest):
    """Update candidate review/shortlist status."""
    valid_statuses = ["shortlisted", "in_review", "rejected", "unreviewed"]
    if req.status.lower() not in valid_statuses:
        raise HTTPException(status_code=400, detail="Invalid status")
        
    success = update_candidate_status(candidate_id, req.status.lower())
    return {"status": "success", "candidate_id": candidate_id, "new_status": req.status.lower()}

@app.delete("/api/candidates/{candidate_id}")
async def remove_candidate(candidate_id: int):
    """Delete a candidate from database and remove their uploaded file if applicable."""
    success = delete_candidate(candidate_id)
    if not success:
        raise HTTPException(status_code=404, detail="Candidate not found")
    return {"status": "success", "message": f"Candidate {candidate_id} deleted successfully"}

@app.get("/api/file/{candidate_id}")
async def view_candidate_resume(candidate_id: int):
    """
    In-app resume viewer endpoint.
    Streams the resume file inline (PDF, DOCX, TXT, image) so it can be previewed in an iframe or tab.
    """
    c = get_candidate_by_id(candidate_id)
    if not c:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    file_path = c.get("file_path")
    if not file_path or not os.path.exists(file_path):
        filename = c.get("filename")
        if filename:
            cleaned_path = PROJECT_ROOT / "cleaned_resumes" / filename
            if cleaned_path.exists():
                file_path = str(cleaned_path)
            else:
                upload_path = UPLOADS_DIR / filename
                if upload_path.exists():
                    file_path = str(upload_path)
                    
    if not file_path or not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Original resume file not found on disk")
        
    media_type, _ = mimetypes.guess_type(file_path)
    if not media_type:
        if file_path.lower().endswith(".pdf"):
            media_type = "application/pdf"
        elif file_path.lower().endswith(".docx"):
            media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif file_path.lower().endswith(".txt"):
            media_type = "text/plain"
        else:
            media_type = "application/octet-stream"
            
    filename = os.path.basename(file_path)
    return FileResponse(
        path=file_path,
        media_type=media_type,
        headers={"Content-Disposition": f'inline; filename="{filename}"'}
    )

@app.get("/api/file/{candidate_id}/download")
async def download_candidate_resume(candidate_id: int):
    """
    Direct resume download endpoint.
    Sends the resume file with attachment Content-Disposition to trigger immediate browser download.
    """
    c = get_candidate_by_id(candidate_id)
    if not c:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    file_path = c.get("file_path")
    if not file_path or not os.path.exists(file_path):
        filename = c.get("filename")
        if filename:
            cleaned_path = PROJECT_ROOT / "cleaned_resumes" / filename
            if cleaned_path.exists():
                file_path = str(cleaned_path)
            else:
                upload_path = UPLOADS_DIR / filename
                if upload_path.exists():
                    file_path = str(upload_path)
                    
    if not file_path or not os.path.exists(file_path):
        raise HTTPException(status_code=404, detail="Original resume file not found on disk")
        
    media_type, _ = mimetypes.guess_type(file_path)
    if not media_type:
        media_type = "application/octet-stream"
        
    filename = c.get("filename") or os.path.basename(file_path)
    return FileResponse(
        path=file_path,
        media_type=media_type,
        filename=filename,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@app.get("/api/stats")
async def fetch_statistics(active_jd_id: str = "ai-engineer"):
    """Fetch overall system KPIs and category distributions."""
    stats = get_statistics(active_jd_id)
    return stats

@app.post("/api/upload")
async def upload_resumes(
    background_tasks: BackgroundTasks,
    files: List[UploadFile] = File(...),
    active_jd_id: str = Form("ai-engineer")
):
    """
    Batch resume upload endpoint.
    Accepts 1, 10, 100+ resumes (PDF, DOCX, TXT).
    Processes extraction, classification, JD scoring, and adds to database.
    Automatically evaluates every uploaded resume using Azure AI.
    """
    jd = get_jd_by_id(active_jd_id)
    results = []
    errors = []
    has_unevaluated = False

    # For batches up to 5 resumes, evaluate with AI synchronously so the user sees results immediately
    evaluate_sync = len(files) <= 5

    for file in files:
        filename = file.filename
        dest_path = UPLOADS_DIR / filename
        
        try:
            # Save uploaded file
            with open(dest_path, "wb") as buffer:
                shutil.copyfileobj(file.file, buffer)
                
            # Extract text
            raw_text = extract_text_from_file(str(dest_path))
            if not raw_text.strip():
                errors.append({"filename": filename, "error": "Unreadable file: no text could be extracted"})
                continue
                
            # Parse resume fields (Passout year explicitly extracted, including 2027)
            parsed = parse_resume(filename, raw_text)
            parsed["file_path"] = str(dest_path.resolve())
            
            # Default local scoring
            jd_scores = score_all_default_jds(parsed)
            parsed["jd_scores"] = jd_scores

            # Automatic AI Evaluation while uploading
            ai_evaluated_done = False
            if evaluate_sync:
                ai_data = analyze_resume_with_ai(raw_text, filename=filename, active_jd_id=active_jd_id)
                if ai_data.get("ai_evaluated"):
                    ai_evaluated_done = True
                    if ai_data.get("name") and ai_data["name"] != "Candidate":
                        parsed["name"] = ai_data["name"]
                    if ai_data.get("email"):
                        parsed["email"] = ai_data["email"]
                    if ai_data.get("phone"):
                        parsed["phone"] = ai_data["phone"]
                    if ai_data.get("passout_year"):
                        parsed["passout_year"] = ai_data["passout_year"]
                    if ai_data.get("cgpa"):
                        parsed["cgpa"] = ai_data["cgpa"]
                    if ai_data.get("degree"):
                        parsed["degree"] = ai_data["degree"]
                    if ai_data.get("primary_category"):
                        parsed["primary_category"] = ai_data["primary_category"]
                    if ai_data.get("secondary_categories"):
                        parsed["secondary_categories"] = ai_data["secondary_categories"]
                    if ai_data.get("skills"):
                        parsed["skills"] = ai_data["skills"]
                    if ai_data.get("projects"):
                        parsed["projects"] = ai_data["projects"]
                    parsed["ai_summary"] = ai_data.get("ai_summary", "")
                    parsed["ai_evaluated"] = 1
                    
                    # Use AI-computed JD scores
                    ai_scores = ai_data.get("scores", {})
                    for k, v in ai_scores.items():
                        jd_scores[k] = float(v)

            if not ai_evaluated_done:
                # Fallback to local rule-based classifier while queued for background AI evaluation
                classification = classify_resume(parsed["skills"], parsed["projects"], parsed["raw_text"])
                parsed["primary_category"] = classification["primary_category"]
                parsed["secondary_categories"] = classification["secondary_categories"]
                parsed["ai_summary"] = ""
                parsed["ai_evaluated"] = 0
                has_unevaluated = True
            
            # Insert into database
            cand_id = upsert_candidate(parsed)
            
            # Active JD score
            active_score = jd_scores.get(active_jd_id, 0.0)
            
            results.append({
                "id": cand_id,
                "filename": filename,
                "name": parsed["name"],
                "passout_year": parsed["passout_year"],
                "cgpa": parsed["cgpa"],
                "primary_category": parsed["primary_category"],
                "secondary_categories": parsed["secondary_categories"],
                "score": active_score,
                "ai_evaluated": parsed.get("ai_evaluated", 0) == 1,
                "ai_summary": parsed.get("ai_summary", ""),
                "status": "success"
            })
        except Exception as e:
            errors.append({"filename": filename, "error": str(e)})

    # If any uploaded resume needs background AI evaluation, launch the background evaluator immediately
    if has_unevaluated:
        evaluator = BatchEvaluator.get_instance()
        background_tasks.add_task(evaluator.start, active_jd_id, False, 6)

    return {
        "total_uploaded": len(files),
        "successful_count": len(results),
        "error_count": len(errors),
        "results": results,
        "errors": errors
    }

@app.post("/api/candidates/{candidate_id}/ai-evaluate")
async def evaluate_candidate_with_azure_ai(candidate_id: int, jd_id: str = "ai-engineer"):
    """
    Triggers Azure OpenAI (gpt-6-luna) to re-read resume, auto-classify type,
    and compute fit scores against JDs.
    """
    c = get_candidate_by_id(candidate_id)
    if not c:
        raise HTTPException(status_code=404, detail="Candidate not found")
        
    raw_text = c.get("raw_text") or ""
    if not raw_text.strip():
        file_path = c.get("file_path")
        if file_path and os.path.exists(file_path):
            raw_text = extract_text_from_file(file_path)
            
    if not raw_text.strip():
        raise HTTPException(status_code=400, detail="No readable resume text available")
        
    ai_result = analyze_resume_with_ai(raw_text, filename=c.get("filename", ""), active_jd_id=jd_id)
    if not ai_result.get("ai_evaluated"):
        raise HTTPException(status_code=502, detail=f"Azure AI evaluation failed: {ai_result.get('error')}")
        
    update_candidate_ai_evaluation(candidate_id, ai_result)
    updated = get_candidate_by_id(candidate_id, active_jd_id=jd_id)
    jd = get_jd_by_id(jd_id)
    jd_analysis = score_resume_against_jd(updated, jd)
    
    score_val = updated.get("active_score")
    if score_val is not None and score_val > 0:
        jd_analysis["score"] = score_val
    else:
        updated["active_score"] = jd_analysis.get("score", 0.0)
    
    return {
        "status": "success",
        "candidate": updated,
        "jd_analysis": jd_analysis,
        "ai_result": ai_result
    }

@app.get("/api/batch-evaluate/status")
async def get_batch_eval_status():
    """Returns current status and progress of background AI batch evaluation."""
    evaluator = BatchEvaluator.get_instance()
    return evaluator.get_status()

@app.post("/api/batch-evaluate/start")
async def start_batch_eval(
    active_jd_id: str = "ai-engineer",
    reevaluate_all: bool = False,
    workers: int = 6
):
    """Triggers background Azure OpenAI batch evaluation across candidates."""
    evaluator = BatchEvaluator.get_instance()
    return evaluator.start(active_jd_id=active_jd_id, reevaluate_all=reevaluate_all, max_workers=workers)

@app.post("/api/batch-evaluate/stop")
async def stop_batch_eval():
    """Stops background Azure OpenAI batch evaluation."""
    evaluator = BatchEvaluator.get_instance()
    return evaluator.stop()


def clean_text_for_export(val: Any) -> str:
    """Normalize internal newlines, nulls, and whitespace for clean spreadsheet rendering."""
    if val is None:
        return ""
    text = str(val).strip()
    text = text.replace("\r\n", " ").replace("\n", " ").replace("\r", " ")
    return " ".join(text.split())

def format_projects_for_export(raw_projects: Any) -> str:
    """Format projects list into clean readable single-line format."""
    if not raw_projects:
        return ""
    if isinstance(raw_projects, str):
        try:
            raw_projects = json.loads(raw_projects)
        except Exception:
            return clean_text_for_export(raw_projects)
    
    formatted = []
    for p in raw_projects:
        if isinstance(p, dict):
            title = clean_text_for_export(p.get("title", ""))
            desc = clean_text_for_export(p.get("description", ""))
            if title and desc:
                formatted.append(f"{title}: {desc}")
            elif title:
                formatted.append(title)
            elif desc:
                formatted.append(desc)
        elif isinstance(p, str):
            cleaned = clean_text_for_export(p)
            if cleaned:
                formatted.append(cleaned)
    return " | ".join(formatted)

@app.get("/api/export")
@app.get("/api/export/excel")
async def export_candidates_csv(
    request: Request,
    search: str = Query(""),
    category: str = Query(""),
    passout_year: Optional[str] = Query(None),
    min_cgpa: Optional[str] = Query(None),
    min_score: Optional[str] = Query(None),
    status: str = Query(""),
    active_jd_id: str = Query("ai-engineer"),
    sort_by: str = Query("score"),
    sort_order: str = Query("desc"),
    format: str = Query("csv")
):
    """
    Export filtered candidate dataset with comprehensive details.
    Supports:
    - CSV format (with UTF-8 BOM, downloadable resume links, and RFC-4180 single-line clean fields)
    - Excel .xlsx format (with native AutoFilter enabled on all columns, styled headers, and clickable hyperlinks)
    """
    py_val = None
    if passout_year and str(passout_year).strip():
        try:
            py_val = int(passout_year)
        except ValueError:
            pass

    cgpa_val = None
    if min_cgpa and str(min_cgpa).strip():
        try:
            cgpa_val = float(min_cgpa)
        except ValueError:
            pass

    score_val = None
    if min_score and str(min_score).strip():
        try:
            score_val = float(min_score)
        except ValueError:
            pass

    # Check if format requested is Excel (.xlsx)
    is_excel = format.lower() in ("xlsx", "excel") or request.url.path.endswith("/excel")

    data = get_candidates(
        search=search,
        category=category,
        passout_year=py_val,
        min_cgpa=cgpa_val,
        min_score=score_val,
        status=status,
        active_jd_id=active_jd_id,
        sort_by=sort_by,
        sort_order=sort_order,
        page=1,
        page_size=100000  # Return all matching candidates
    )
    
    candidates = data["candidates"]

    # Retrieve Active JD Details
    active_jd = get_jd_by_id(active_jd_id)
    jd_title = active_jd.get("title", active_jd_id) if active_jd else active_jd_id

    # Compute base URL for clickable resume download links
    base_url = str(request.base_url).rstrip("/")

    headers = [
        "Candidate ID",
        "Candidate Name",
        "Email Address",
        "Phone Number",
        "Resume Download Link",
        "Resume Action (Excel Click)",
        "Resume View Link",
        "Passout Batch",
        "CGPA",
        "Degree / Qualification",
        "Primary Domain",
        "Secondary Domains",
        "Top Skills",
        "All Technical Skills",
        "Target Job Description",
        "JD Fit Match (%)",
        "AI Engineer Fit (%)",
        "Full Stack Java Fit (%)",
        "AR Caller Fit (%)",
        "US BD Sales Fit (%)",
        "Review Status",
        "Azure AI Evaluated",
        "AI Executive Summary",
        "Key Projects & Highlights",
        "GitHub Profile",
        "LinkedIn Profile",
        "Original Filename",
        "Record Created At"
    ]

    rows = []
    for c in candidates:
        cand_id = c["id"]
        download_url = f"{base_url}/api/file/{cand_id}/download"
        view_url = f"{base_url}/api/file/{cand_id}"
        excel_hyperlink = f'=HYPERLINK("{download_url}", "Download Resume")'

        skills_list = c.get("skills") or []
        top_skills = ", ".join(skills_list[:10])
        all_skills = ", ".join(skills_list)
        sec_cats = ", ".join(c.get("secondary_categories") or [])
        projects_str = format_projects_for_export(c.get("projects"))
        ai_summary_str = clean_text_for_export(c.get("ai_summary", ""))

        row = [
            cand_id,
            clean_text_for_export(c.get("name") or "Candidate"),
            clean_text_for_export(c.get("email") or ""),
            clean_text_for_export(c.get("phone") or ""),
            download_url,
            excel_hyperlink,
            view_url,
            c.get("passout_year") or "",
            (round(c["cgpa"] / 10.0, 2) if c["cgpa"] > 10 else round(c["cgpa"], 2)) if c.get("cgpa") is not None else "-",
            clean_text_for_export(c.get("degree") or ""),
            clean_text_for_export(c.get("primary_category") or "Other"),
            sec_cats,
            top_skills,
            all_skills,
            jd_title,
            round(float(c.get("active_score") or 0), 1),
            round(float(c.get("score_ai_engineer") or 0), 1),
            round(float(c.get("score_full_stack_java") or 0), 1),
            round(float(c.get("score_ar_caller") or 0), 1),
            round(float(c.get("score_us_bd_sales") or 0), 1),
            (c.get("status") or "unreviewed").capitalize(),
            "Yes" if c.get("ai_evaluated") else "No",
            ai_summary_str,
            projects_str,
            clean_text_for_export(c.get("github") or ""),
            clean_text_for_export(c.get("linkedin") or ""),
            clean_text_for_export(c.get("filename") or ""),
            clean_text_for_export(c.get("created_at") or "")
        ]
        rows.append(row)

    file_slug = active_jd_id.replace(" ", "_").lower()

    if is_excel and HAS_OPENPYXL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Candidate Roster"
        ws.views.sheetView[0].showGridLines = True

        # Append headers
        ws.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        ws.row_dimensions[1].height = 28

        link_font = Font(name="Calibri", size=10, color="004AAC", underline="single")

        # Append data rows
        for r_idx, row_data in enumerate(rows, start=2):
            ws.append(row_data)
            # Add native Excel clickable hyperlink to Resume Download Link (column 5)
            link_cell = ws.cell(row=r_idx, column=5)
            link_cell.hyperlink = row_data[4]
            link_cell.font = link_font

            # View link (column 7)
            view_cell = ws.cell(row=r_idx, column=7)
            view_cell.hyperlink = row_data[6]
            view_cell.font = link_font

        # CRITICAL: Enable AutoFilter across ALL COLUMNS!
        max_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{max_col_letter}{len(rows) + 1}"

        # Freeze top row so headers remain visible when scrolling
        ws.freeze_panes = "A2"

        # Auto-adjust column widths
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in list(col)[:50]:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

        output_buf = io.BytesIO()
        wb.save(output_buf)
        output_buf.seek(0)
        return StreamingResponse(
            output_buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="candidates_export_{file_slug}.xlsx"'}
        )

    # Standard CSV Export with UTF-8 BOM
    output = io.StringIO()
    output.write("\ufeff")  # UTF-8 BOM for Microsoft Excel auto-encoding
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)

    output.seek(0)
    return StreamingResponse(
        io.BytesIO(output.getvalue().encode("utf-8")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="candidates_export_{file_slug}.csv"'}
    )


@app.get("/{full_path:path}", response_class=HTMLResponse)
async def serve_spa_routes(full_path: str):
    """Serve React frontend SPA for client-side routes (/candidates, /jobs, /upload, etc.)."""
    if full_path.startswith("api/") or full_path.startswith("assets/") or full_path.startswith("static/"):
        raise HTTPException(status_code=404, detail="Resource not found")
        
    react_index = FRONTEND_DIST / "index.html"
    if react_index.exists():
        with open(react_index, "r", encoding="utf-8") as f:
            return HTMLResponse(f.read())
            
    return HTMLResponse("<h1>Loading React Platform...</h1>", status_code=200)

